from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import anyio
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.companies.normalization import normalize_employer_name, normalize_https_url
from app.db.models import ImmigrationDatasetImport, LcaEmployerYearlyStat

_EMPLOYER_COLUMNS = ("EMPLOYER_NAME", "EMPLOYER_BUSINESS_DBA")
_STATUS_COLUMNS = ("CASE_STATUS", "STATUS")
_WORKER_COLUMNS = ("TOTAL_WORKER_POSITIONS", "WORKSITE_WORKERS", "WORKERS")
_OCCUPATION_COLUMNS = ("SOC_TITLE", "JOB_TITLE", "OCCUPATIONAL_TITLE")


class LcaImportError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass
class EmployerAggregate:
    legal_name: str
    normalized_name: str
    certified_cases: int = 0
    certified_workers: int = 0
    denied_cases: int = 0
    withdrawn_cases: int = 0
    occupations: Counter[str] = field(default_factory=Counter)


@dataclass(frozen=True)
class ImportSummary:
    import_id: str
    fiscal_year: int
    rows_read: int
    rows_accepted: int
    employers: int
    cached: bool
    source_sha256: str


class LcaImporter:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def import_file(
        self,
        file_path: Path,
        fiscal_year: int,
        source_url: str,
        record_layout_version: str,
    ) -> ImportSummary:
        if not 2000 <= fiscal_year <= datetime.now(UTC).year + 1:
            raise LcaImportError(
                "INVALID_FISCAL_YEAR", "Fiscal year is outside the supported range."
            )
        if not await anyio.Path(file_path).is_file():
            raise LcaImportError("FILE_NOT_FOUND", f"File does not exist: {file_path}")
        normalized_source_url = normalize_https_url(source_url)
        source_sha256 = await anyio.to_thread.run_sync(_sha256, file_path)
        import_row = await self.session.scalar(
            select(ImmigrationDatasetImport).where(
                ImmigrationDatasetImport.dataset_type == "dol_lca",
                ImmigrationDatasetImport.fiscal_year == fiscal_year,
                ImmigrationDatasetImport.source_sha256 == source_sha256,
            )
        )
        if import_row is not None and import_row.status == "succeeded":
            employer_count = len(
                (
                    await self.session.scalars(
                        select(LcaEmployerYearlyStat.id).where(
                            LcaEmployerYearlyStat.dataset_import_id == import_row.id
                        )
                    )
                ).all()
            )
            return _summary(import_row, employer_count, cached=True)

        if import_row is None:
            import_row = ImmigrationDatasetImport(
                dataset_type="dol_lca",
                fiscal_year=fiscal_year,
                source_url=normalized_source_url,
                source_sha256=source_sha256,
                record_layout_version=record_layout_version,
                status="running",
            )
            self.session.add(import_row)
        else:
            import_row.status = "running"
            import_row.error_code = None
            import_row.rows_read = 0
            import_row.rows_accepted = 0
            import_row.source_url = normalized_source_url
            import_row.record_layout_version = record_layout_version
        await self.session.commit()

        try:
            rows_read, rows_accepted, aggregates = await anyio.to_thread.run_sync(
                aggregate_lca_file, file_path
            )
            for aggregate in aggregates.values():
                self.session.add(
                    LcaEmployerYearlyStat(
                        dataset_import_id=import_row.id,
                        fiscal_year=fiscal_year,
                        legal_employer_name=aggregate.legal_name,
                        normalized_legal_employer_name=aggregate.normalized_name,
                        certified_cases=aggregate.certified_cases,
                        certified_workers=aggregate.certified_workers,
                        denied_cases=aggregate.denied_cases,
                        withdrawn_cases=aggregate.withdrawn_cases,
                        top_occupations=[
                            {"name": name, "cases": count}
                            for name, count in aggregate.occupations.most_common(5)
                        ],
                    )
                )
            import_row.rows_read = rows_read
            import_row.rows_accepted = rows_accepted
            import_row.status = "succeeded"
            import_row.imported_at = datetime.now(UTC)
            await self.session.commit()
        except Exception as exc:
            await self.session.rollback()
            import_row = await self.session.get(ImmigrationDatasetImport, import_row.id)
            if import_row is not None:
                import_row.status = "failed"
                import_row.error_code = (
                    exc.code if isinstance(exc, LcaImportError) else "IMPORT_FAILED"
                )
                await self.session.commit()
            if isinstance(exc, LcaImportError):
                raise
            raise LcaImportError(
                "IMPORT_FAILED", "The H-1B sponsorship file could not be imported."
            ) from exc

        return _summary(import_row, len(aggregates), cached=False)


def aggregate_lca_file(file_path: Path) -> tuple[int, int, dict[str, EmployerAggregate]]:
    aggregates: dict[str, EmployerAggregate] = {}
    rows_read = 0
    rows_accepted = 0
    for record in _iter_records(file_path):
        rows_read += 1
        legal_name = _string_value(record.get("employer"))
        status = _string_value(record.get("status")).upper()
        normalized_name = normalize_employer_name(legal_name)
        if not normalized_name or not status:
            continue
        aggregate = aggregates.setdefault(
            normalized_name,
            EmployerAggregate(legal_name=legal_name, normalized_name=normalized_name),
        )
        workers = _positive_int(record.get("workers"), default=1)
        if status.startswith("CERTIFIED"):
            aggregate.certified_cases += 1
            aggregate.certified_workers += workers
        if "DENIED" in status:
            aggregate.denied_cases += 1
        if "WITHDRAWN" in status:
            aggregate.withdrawn_cases += 1
        occupation = _string_value(record.get("occupation"))
        if occupation:
            aggregate.occupations[occupation] += 1
        rows_accepted += 1
    return rows_read, rows_accepted, aggregates


def _iter_records(file_path: Path) -> Iterator[dict[str, Any]]:
    suffix = file_path.suffix.casefold()
    if suffix == ".csv":
        yield from _iter_csv(file_path)
        return
    if suffix in {".xlsx", ".xlsm"}:
        yield from _iter_xlsx(file_path)
        return
    raise LcaImportError("UNSUPPORTED_FILE_TYPE", "Use an official CSV or XLSX disclosure file.")


def _iter_csv(file_path: Path) -> Iterator[dict[str, Any]]:
    try:
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            headers = next(reader, None)
            if headers is None:
                raise LcaImportError("EMPTY_FILE", "The disclosure file is empty.")
            layout = _resolve_layout(headers)
            for row in reader:
                yield _canonical_record(row, layout)
    except UnicodeDecodeError as exc:
        raise LcaImportError("INVALID_ENCODING", "CSV files must use UTF-8 encoding.") from exc


def _iter_xlsx(file_path: Path) -> Iterator[dict[str, Any]]:
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        raise LcaImportError("INVALID_XLSX", "The XLSX workbook could not be opened.") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise LcaImportError("EMPTY_FILE", "The disclosure workbook has no active sheet.")
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise LcaImportError("EMPTY_FILE", "The disclosure file is empty.")
        layout = _resolve_layout([_string_value(value) for value in headers])
        for row in rows:
            yield _canonical_record(list(row), layout)
    finally:
        workbook.close()


def _resolve_layout(headers: list[str]) -> dict[str, int | None]:
    normalized_headers = [_normalize_header(header) for header in headers]
    employer = _first_index(normalized_headers, _EMPLOYER_COLUMNS)
    status = _first_index(normalized_headers, _STATUS_COLUMNS)
    if employer is None or status is None:
        raise LcaImportError(
            "MISSING_REQUIRED_COLUMNS",
            "The file must contain employer-name and case-status columns.",
        )
    return {
        "employer": employer,
        "status": status,
        "workers": _first_index(normalized_headers, _WORKER_COLUMNS),
        "occupation": _first_index(normalized_headers, _OCCUPATION_COLUMNS),
    }


def _canonical_record(
    row: list[Any] | tuple[Any, ...], layout: dict[str, int | None]
) -> dict[str, Any]:
    return {
        name: row[index] if index is not None and index < len(row) else None
        for name, index in layout.items()
    }


def _normalize_header(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", value.strip().upper()).strip("_")


def _first_index(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for candidate in candidates:
        try:
            return headers.index(candidate)
        except ValueError:
            continue
    return None


def _string_value(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _positive_int(value: Any, default: int) -> int:
    try:
        parsed = int(float(str(value)))
    except (TypeError, ValueError):
        return default
    return max(parsed, 0)


def _sha256(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _summary(
    import_row: ImmigrationDatasetImport, employer_count: int, *, cached: bool
) -> ImportSummary:
    return ImportSummary(
        import_id=str(import_row.id),
        fiscal_year=import_row.fiscal_year,
        rows_read=import_row.rows_read,
        rows_accepted=import_row.rows_accepted,
        employers=employer_count,
        cached=cached,
        source_sha256=import_row.source_sha256,
    )
